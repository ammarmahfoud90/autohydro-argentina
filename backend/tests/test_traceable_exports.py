"""End-to-end regressions: recomputation, traceability and export consistency."""

from io import BytesIO

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
import pytest

from app.services.calculation import run_calculation
from app.services.calculation_trace import prepare_report_data
from app.services.report_service import MemoriaCalculoGenerator
from app.services.docx_service import MemoriaCalculoDocxGenerator
from app.services.excel_service import ExcelReportGenerator
from app.services.hydraulic_report_service import CulvertReportGenerator
from app.services.culvert_service import calculate_culvert, get_reference_data
from app.services.proyecto_report_service import ProyectoReportGenerator


def calculation():
    return run_calculation(dict(locality_id="amgr", return_period=25,
        duration_min=60, area_km2=1, length_km=1, slope=0.01,
        method="rational", runoff_coeff=0.6, tc_formulas=["temez"]))


def test_client_edited_values_are_recomputed_from_original_inputs():
    original = calculation()
    tampered = {**original, "peak_flow_m3s": 999999, "duration_min": 60}
    restored = prepare_report_data(tampered)
    assert restored["peak_flow_m3s"] == original["peak_flow_m3s"]
    assert restored["duration_min"] == original["effective_duration_min"]
    assert restored["length_km"] == 1
    assert restored["slope"] == 0.01
    assert restored["risk_level"] == "no_evaluado"


def test_legacy_report_is_explicitly_unreproducible():
    legacy = {"peak_flow_m3s": 1, "risk_level": "bajo", "return_period": 25}
    result = prepare_report_data(prepare_report_data(legacy))
    assert result["validation_status"] == "legacy_not_reproducible"
    assert result["risk_level"] == "no_evaluado"
    assert len(result["warnings"]) == 1


@pytest.mark.parametrize("format", ["pdf", "docx", "xlsx"])
def test_exports_preserve_inputs_source_warnings_and_actual_duration(format):
    data = calculation()
    if format == "pdf":
        buffer = MemoriaCalculoGenerator().generate(data)
        text = " ".join(page.extract_text() for page in PdfReader(buffer).pages)
    elif format == "docx":
        buffer = MemoriaCalculoDocxGenerator().generate(data)
        document = Document(buffer)
        text = " ".join([p.text for p in document.paragraphs] +
                        [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    else:
        buffer = ExcelReportGenerator().generate(data)
        workbook = load_workbook(buffer, data_only=True)
        text = " ".join(str(cell.value) for sheet in workbook for row in sheet for cell in row if cell.value is not None)
    assert "length_km" in text and "slope" in text
    assert data["calculation_version"] in text
    assert data["idf_data_sha256"] in text.replace("\n", "")
    assert "NO EVALUADO" in text
    assert "duración" in text.lower() and "ajustada" in text.lower()
    assert str(data["effective_duration_min"]) in text


def test_culvert_pdf_preserves_tailwater_and_model_assumptions():
    data = calculate_culvert(dict(design_flow_m3s=0.5, culvert_type="circular",
        material="hormigon", length_m=10, slope=0.005,
        inlet_type="sin_alas", headwater_max_m=2, tailwater_m=0.3))
    buffer = CulvertReportGenerator("Caso de referencia", "Argentina").generate(data)
    text = " ".join(page.extract_text() for page in PdfReader(buffer).pages)
    assert "TW" in text and "0.30" in text
    assert "tabla A.1" in text
    assert "predimensionamiento" in text.lower()


def test_unsupported_inlet_geometry_is_not_offered_as_a_valid_reference():
    assert "redondeado" not in {entry["id"] for entry in get_reference_data()["inlet_types"]}
    with pytest.raises(ValueError, match="geometría"):
        calculate_culvert(dict(design_flow_m3s=0.5, culvert_type="circular",
            length_m=10, slope=0.005, inlet_type="redondeado", headwater_max_m=2))


def test_consolidated_project_cannot_export_a_stale_flow_and_capacity_check():
    data = calculation()
    project = {"paso4": {"result": data, "q_pico_m3s": 999}, "paso5": {"verifica": True}}
    with pytest.raises(ValueError, match="recalcular"):
        ProyectoReportGenerator().generate(project)


@pytest.mark.parametrize("format,endpoint", [
    ("pdf", "/api/report"), ("docx", "/api/report/docx"), ("xlsx", "/api/report/excel"),
])
def test_http_calculation_to_report_keeps_reproducible_data(format, endpoint):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.api.routes import hydrology, report
    from app.core.limiter import limiter

    app = FastAPI()
    app.state.limiter = limiter
    app.include_router(hydrology.router, prefix="/api")
    app.include_router(report.router, prefix="/api")
    client = TestClient(app)
    data = calculation()
    response = client.post("/api/calculate", json=data["calculation_inputs"])
    assert response.status_code == 200
    serialized = response.json()
    assert serialized["calculation_inputs"]["length_km"] == 1
    assert serialized["duration_min"] == serialized["effective_duration_min"]
    exported = client.post(endpoint, json={"calculationData": serialized, "fetchAISections": False})
    assert exported.status_code == 200
    assert len(exported.content) > 5000
    disposition = exported.headers["content-disposition"]
    assert format in disposition
    assert disposition.isascii()
    assert 'filename="memoria_calculo_Proyecto Hidrologico.' in disposition
    assert "filename*=UTF-8''memoria_calculo_Proyecto%20Hidrol%C3%B3gico." in disposition
